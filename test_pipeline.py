import unittest
from pathlib import Path
from tempfile import TemporaryDirectory
from openpyxl import load_workbook
from pipeline import export, ingest, safe_cell

class PipelineTests(unittest.TestCase):
    def test_alignment_is_one_to_one(self):
        rows = ingest()
        self.assertEqual(len(rows), 6)
        self.assertEqual(len({r['video'] for r in rows}), 6)
        self.assertTrue(all(r['transcript'] and 'T' in r['recorded_date'] for r in rows))

    def test_spreadsheet_formula_escape(self):
        self.assertEqual(safe_cell('=HYPERLINK("bad")'), '\'=HYPERLINK("bad")')

    def test_workbook_preserves_candidates_and_selection(self):
        e = dict(video=Path('example.mp4'), name='Example', source_sheet='Sheet1', source_row=3, recorded_date='2026-09-11')
        candidates = [dict(id=f'fst{i:08d}', heading=f'Heading {i}', facet='Topical', score=1/(i+1)) for i in range(30)]
        result = dict(description=dict(description='Example description', visual_observations=[], uncertainties=[]),
                      selection=dict(headings=[dict(id=r['id'], rationale='Relevant') for r in candidates[:5]], keywords=['a','b','c','d','e']),
                      candidates=candidates, model='test', technical=dict(duration=60), usage=[], sample_timestamps=[])
        with TemporaryDirectory() as tmp:
            path = Path(tmp)/'result.xlsx'
            export([(e, result)], path)
            wb = load_workbook(path)
            self.assertEqual(wb['Video metadata'].max_row, 2)
            self.assertEqual(wb['Top 30 FAST candidates'].max_row, 31)
            self.assertEqual(wb['Video metadata']['D2'].value, 'Heading 0 (fst00000000)')
            wb.close()

if __name__ == '__main__':
    unittest.main()
